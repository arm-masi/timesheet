import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.schemas.attendance import MonthlySummary


def _status_label(d) -> str:
    """Get human-readable status for a daily row."""
    if d.has_justification:
        tipo = "Ferie" if d.justification_type == "ferie" else "Permesso"
        stato = {
            "approvato": "Approvato",
            "rifiutato": "Rifiutato",
            "in_attesa": "In attesa",
        }.get(d.justification_status or "", d.justification_status or "")
        return f"{tipo} ({stato})"
    if d.is_missing:
        return "Da giustificare"
    if d.clock_in and d.clock_out:
        return "Presente"
    return ""


def export_to_csv(summaries: list[MonthlySummary]) -> str:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Dipendente", "Data", "Entrata", "Uscita",
        "Ore Lavorate", "Deficit", "Stato"
    ])
    for s in summaries:
        for d in s.daily_details:
            writer.writerow([
                s.user_name,
                d.date,
                d.clock_in or "",
                d.clock_out or "",
                d.worked_hours,
                d.deficit_hours if d.deficit_hours > 0 else "",
                _status_label(d),
            ])
        # Summary row per user
        writer.writerow([
            s.user_name,
            "TOTALE",
            "",
            "",
            s.total_hours,
            "",
            f"Lavorati: {s.total_worked_days} | Ferie: {s.total_ferie_days} | Permessi: {s.total_permesso_days} | Non giust.: {s.total_missing_days}",
        ])
        writer.writerow([])  # blank separator
    return output.getvalue()


def export_to_excel(summaries: list[MonthlySummary]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dettaglio Paghe"

    headers = [
        "Dipendente", "Data", "Entrata", "Uscita",
        "Ore Lavorate", "Deficit", "Stato"
    ]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    summary_fill = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    summary_font = Font(bold=True)
    missing_font = Font(color="E53E3E")

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for s in summaries:
        for d in s.daily_details:
            status = _status_label(d)
            row = [
                s.user_name,
                d.date,
                d.clock_in or "",
                d.clock_out or "",
                d.worked_hours,
                d.deficit_hours if d.deficit_hours > 0 else "",
                status,
            ]
            ws.append(row)
            if d.is_missing:
                row_idx = ws.max_row
                for col in range(1, 8):
                    ws.cell(row=row_idx, column=col).font = missing_font

        # Summary row
        summary_text = f"Lavorati: {s.total_worked_days} | Ferie: {s.total_ferie_days} | Permessi: {s.total_permesso_days} | Non giust.: {s.total_missing_days}"
        ws.append([s.user_name, "TOTALE", "", "", s.total_hours, "", summary_text])
        row_idx = ws.max_row
        for col in range(1, 8):
            ws.cell(row=row_idx, column=col).fill = summary_fill
            ws.cell(row=row_idx, column=col).font = summary_font

        ws.append([])  # blank separator

    # Auto-size columns
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
